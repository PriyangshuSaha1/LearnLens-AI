"""
Document Loader Module for LearnLens AI.
Handles loading and text extraction from PDF, PPTX, DOCX, and TXT files.
"""

import os
import tempfile
from typing import List

from langchain_core.documents import Document


def _extract_pdf(file_path: str, source_name: str) -> List[Document]:
    """Extract text from a PDF file, one Document per page."""
    from PyPDF2 import PdfReader

    docs = []
    reader = PdfReader(file_path)
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": source_name,
                        "page": i + 1,
                        "total_pages": len(reader.pages),
                        "file_type": "pdf",
                    },
                )
            )
    return docs


def _extract_docx(file_path: str, source_name: str) -> List[Document]:
    """Extract text from a DOCX file."""
    import docx

    doc = docx.Document(file_path)
    full_text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    if full_text.strip():
        return [
            Document(
                page_content=full_text,
                metadata={"source": source_name, "file_type": "docx"},
            )
        ]
    return []


def _extract_pptx(file_path: str, source_name: str) -> List[Document]:
    """Extract text from a PPTX file, one Document per slide."""
    from pptx import Presentation

    docs = []
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
        slide_text = "\n".join(texts)
        if slide_text.strip():
            docs.append(
                Document(
                    page_content=slide_text,
                    metadata={
                        "source": source_name,
                        "slide": i + 1,
                        "total_slides": len(prs.slides),
                        "file_type": "pptx",
                    },
                )
            )
    return docs


def _extract_txt(file_path: str, source_name: str) -> List[Document]:
    """Extract text from a plain text file."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    if text.strip():
        return [
            Document(
                page_content=text,
                metadata={"source": source_name, "file_type": "txt"},
            )
        ]
    return []


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".txt": _extract_txt,
}


def load_document(uploaded_file) -> List[Document]:
    """
    Load a single Streamlit UploadedFile and return a list of LangChain Documents.

    Args:
        uploaded_file: A Streamlit UploadedFile object.

    Returns:
        A list of Document objects with page_content and metadata.

    Raises:
        ValueError: If the file type is not supported.
    """
    file_name = uploaded_file.name
    _, ext = os.path.splitext(file_name)
    ext = ext.lower()

    if ext not in _EXTRACTORS:
        raise ValueError(
            f"Unsupported file type: '{ext}'. Supported: {list(_EXTRACTORS.keys())}"
        )

    # Write the uploaded bytes to a temp file for processing
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        tmp.write(uploaded_file.getbuffer())
        tmp_path = tmp.name

    try:
        documents = _EXTRACTORS[ext](tmp_path, file_name)
    finally:
        os.unlink(tmp_path)

    return documents


def load_multiple_documents(uploaded_files) -> List[Document]:
    """
    Load multiple Streamlit UploadedFile objects.

    Args:
        uploaded_files: A list of Streamlit UploadedFile objects.

    Returns:
        A combined list of Document objects from all files.
    """
    all_docs = []
    for uf in uploaded_files:
        try:
            docs = load_document(uf)
            all_docs.extend(docs)
        except Exception as e:
            print(f"Warning: Failed to load '{uf.name}': {e}")
    return all_docs
